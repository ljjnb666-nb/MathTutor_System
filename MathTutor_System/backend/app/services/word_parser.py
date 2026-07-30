"""
Word 试卷解析服务：从 .docx 提取文本，经 LLM 解析为结构化题目 JSON。
支持 Word→PDF→按页识图流程：含图试卷可先转 PDF，每页渲染成图后由视觉模型识别题目。
解析后先做「同题号合并」再做「小节续写合并」；小节合并采用白名单 + 通用【…】模式，
遇到综合与探究/综合与实践等新小节题型一般无需改代码。
"""
import base64
import json
import logging
import re
from io import BytesIO
from typing import Any

from docx import Document
from openai import OpenAI
from PIL import Image

from app.core.config import AI_REQUEST_TIMEOUT, DEEPSEEK_API_KEY, DEEPSEEK_BASE_URL, DEEPSEEK_MODEL
from app.services.exam_question_merge import (
    append_region as _append_region,
    merge_by_section_headers as _merge_by_section_headers,
    merge_same_number_questions as _merge_same_number_questions,
    normalize_question_number as _normalize_question_number,
    normalize_question_type as _normalize_question_type,
    question_merge_key as _question_merge_key,
)

logger = logging.getLogger(__name__)

# 按页识图时的系统提示：与 parse_with_deepseek 输出结构一致；每题可带 image_region 便于只展示该题对应图片
VISION_EXAM_SYSTEM_PROMPT = (
    "You are an exam parser. You will receive an image of one or more exam questions (possibly with diagrams). "
    "Extract every question on this page into a strict JSON array. "
    "IMPORTANT: A question with ONE main number (e.g. '20.' or '20') but multiple sub-sections (e.g. 【探索发现】【抽象定义】【问题解决】【方法应用】, or 【问题背景】【研究条件】【模型构建】【总结反思】, or any 【…】-style section headers) or sub-questions (①②③ or (1)(2)(3)) is ONE question. "
    "Do NOT split it into multiple array items by sub-section or sub-question. Put the entire content (all sections, all ①②③) into a single content string. "
    "Only start a new array item when you see a new main question number (e.g. next is '21.'). "
    "Each item must have: number (int or string, the main question number only), type ('choice' | 'fill' | 'solution'), "
    "content (string, full question text including all sub-parts; if there is a diagram or '如图', include a brief description), "
    "and options (array of strings) only for choice type. "
    "For choice questions: extract EVERY option (A, B, C, D, ...) as a separate string. "
    "Preserve mathematical notation: superscripts (a², a⁴), multiplication dot (·), parentheses. Each option one complete string. "
    "If a question refers to a diagram/figure (e.g. 如图, 示意图), provide image_region: an object with x, y, width, height (0 to 1, normalized). Top-left origin. "
    "image_region must be the bounding box of ONLY the reference figure/diagram (几何图、示意图), NOT including the question text. The inserted image should be just the figure the student needs to look at. "
    "Example: \"image_region\": {\"x\": 0.1, \"y\": 0.2, \"width\": 0.3, \"height\": 0.25}. "
    "For text-only questions that do not need a reference image, omit image_region. "
    "Output ONLY valid JSON array. No Markdown, no code block wrapper."
)


def extract_text_from_docx(file_stream: BytesIO) -> str:
    """
    使用 python-docx 读取上传文件流，遍历段落并拼接为单一文本，过滤空行。
    """
    doc = Document(file_stream)
    lines: list[str] = []
    for para in doc.paragraphs:
        line = (para.text or "").strip()
        if line:
            lines.append(line)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                part = (cell.text or "").strip()
                if part:
                    lines.append(part)
    return "\n".join(lines)


def _clean_json_string(raw: str) -> str:
    """去除 AI 返回的 markdown 代码块标记后再解析。"""
    s = (raw or "").strip()
    # 去掉 ```json ... ``` 或 ``` ... ```
    m = re.search(r"```(?:json)?\s*([\s\S]*?)```", s)
    if m:
        s = m.group(1).strip()
    return s


def parse_with_deepseek(
    text: str,
    api_key: str | None = None,
    base_url: str | None = None,
    model: str | None = None,
    provider: str | None = None,
) -> list[dict[str, Any]]:
    """
    将试卷文本解析为题目 JSON 数组。
    - provider=gemini 时使用 Gemini REST API（直接调用，无需 OpenAI 兼容 Base URL）。
    - 否则使用 OpenAI 兼容接口（DeepSeek 等），base_url/api_key/model 来自设置或 .env。
    """
    provider_lower = (provider or "").strip().lower()
    system_prompt = (
        "You are an exam parser. Convert the input text into a strict JSON array. "
        "Identify questions by main question numbers (1., 2., ..., 20., etc.). "
        "IMPORTANT: A question with one main number but multiple sub-parts (e.g. 【探索发现】【问题解决】【问题背景】【研究条件】【模型构建】【总结反思】 or any 【…】-style section headers, or ①②③, or (1)(2)(3)) is ONE question. "
        "Do NOT split it into multiple array items. Put the entire text of that question (all sections, all sub-questions) into a single content string. "
        "Only create a new array item when the main question number changes (e.g. from 20 to 21). "
        "Detect type: 'choice' (选择题), 'fill' (填空题), or 'solution' (解答题). Multi-part 综合与探究/综合与实践/综合题 are type 'solution'. "
        "If type is 'choice', extract EVERY option (A, B, C, D, ...) as an array of strings; preserve superscripts (a², a⁴), multiplication dot (·), parentheses. "
        "Each item must have: number (int or string), type ('choice' | 'fill' | 'solution'), content (string), and options (array) only for choice. "
        "Output ONLY valid JSON. No Markdown formatting."
    )

    if provider_lower == "gemini":
        key = (api_key or "").strip()
        if not key:
            raise ValueError(
                "未配置 API Key。请在前端「设置」中选择 Gemini 并填写 API Key 后保存，或在 .env 中设置 LLM_API_KEY。"
            )
        model_name = (model or "").strip() or "gemini-1.5-flash"
        try:
            from app.services.llm_service import gemini_rest_text_sync
            full_prompt = f"{system_prompt}\n\nInput:\n{text}"
            raw = gemini_rest_text_sync(full_prompt, api_key=key, model=model_name, temperature=0.2, max_tokens=4096)
        except Exception as e:
            logger.warning("word_parse_failed", extra={"error_type": type(e).__name__, "detail": str(e)}, exc_info=True)
            raise ValueError(f"LLM 解析失败: {str(e)}") from e
    else:
        key = (api_key or "").strip() or DEEPSEEK_API_KEY
        if not key:
            raise ValueError(
                "未配置 API Key。请在前端「设置」中选择模型提供商并填写 API Key 后保存，或在 .env 中设置对应 API Key。"
            )
        url = (base_url or "").strip() or DEEPSEEK_BASE_URL
        client = OpenAI(base_url=url, api_key=key, timeout=float(AI_REQUEST_TIMEOUT))
        model_name = (model or "").strip() or DEEPSEEK_MODEL
        try:
            resp = client.chat.completions.create(
                model=model_name,
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": text},
                ],
                temperature=0.2,
            )
            raw = (resp.choices[0].message.content or "").strip()
        except Exception as e:
            logger.warning("word_parse_failed", extra={"error_type": type(e).__name__, "detail": str(e)}, exc_info=True)
            raise ValueError(f"LLM 解析失败: {str(e)}") from e

    try:
        cleaned = _clean_json_string(raw)
        data = json.loads(cleaned)
        if not isinstance(data, list):
            data = [data] if isinstance(data, dict) else []
        for item in data:
            if isinstance(item, dict):
                item["type"] = _normalize_question_type(
                    item.get("type"), item.get("options")
                )
        data = _merge_same_number_questions(data)
        data = _merge_by_section_headers(data)
        return data
    except json.JSONDecodeError as e:
        logger.warning(
            "word_parse_failed",
            extra={"error_type": "JSONDecodeError", "detail": str(e)},
            exc_info=True,
        )
        raise ValueError(f"LLM 返回内容不是合法 JSON: {e}") from e


def parse_page_image_with_vision(
    page_image_bytes: bytes,
    api_key: str,
    base_url: str,
    model: str,
    provider: str | None = None,
) -> list[dict[str, Any]]:
    """
    对单页试卷图片调用视觉模型，识别该页上的题目并返回与 parse_with_deepseek 同结构的列表。
    provider=gemini 时使用 Gemini REST 多模态接口；否则使用 OpenAI 兼容视觉接口。
    """
    if not (api_key or "").strip():
        raise ValueError("未配置 API Key，无法使用按页识图。请在「设置」中填写并保存。")
    key = (api_key or "").strip()
    b64 = base64.standard_b64encode(page_image_bytes).decode("ascii")
    vision_user_text = (
        "从本页试卷截图中识别所有题目，输出 JSON 数组。"
        "若同一题有多个小节（如【探索发现】【问题解决】【问题背景】【研究条件】【模型构建】【总结反思】等任意【…】形式）或小题①②③，只输出一道题，将全部内容合并到该题的 content 中，不要按小节或小题拆成多道题。"
        "选择题须逐项提取 A、B、C、D 等选项；指数、乘号等数学符号原样保留。"
        "仅当题目引用图（如 如图、示意图）时，为该题提供 image_region（x、y、width、height，0～1 归一化，示例：{\"x\":0.1,\"y\":0.2,\"width\":0.3,\"height\":0.25}）。image_region 只框选参考图/示意图本身，不要包含题干文字。纯文字题不提供 image_region。"
    )
    full_prompt = f"{VISION_EXAM_SYSTEM_PROMPT}\n\n{vision_user_text}"

    provider_lower = (provider or "").strip().lower()
    if provider_lower == "gemini":
        model_name = (model or "").strip() or "gemini-1.5-flash"
        try:
            from app.services.llm_service import gemini_rest_vision_sync
            raw = gemini_rest_vision_sync(
                full_prompt, image_base64=b64, api_key=key, model=model_name,
                mime_type="image/png", temperature=0.2, max_tokens=4096,
            )
        except Exception as e:
            logger.warning("按页识图失败: %s", e)
            raise ValueError(f"按页识图失败: {str(e)}") from e
    else:
        url = (base_url or "").strip() or DEEPSEEK_BASE_URL
        model_name = (model or "").strip() or DEEPSEEK_MODEL
        data_uri = f"data:image/png;base64,{b64}"
        client = OpenAI(base_url=url, api_key=key, timeout=float(AI_REQUEST_TIMEOUT))
        try:
            resp = client.chat.completions.create(
                model=model_name,
                messages=[
                    {"role": "system", "content": VISION_EXAM_SYSTEM_PROMPT},
                    {
                        "role": "user",
                        "content": [
                            {"type": "text", "text": vision_user_text},
                            {"type": "image_url", "image_url": {"url": data_uri}},
                        ],
                    },
                ],
                temperature=0.2,
                max_tokens=4096,
            )
            raw = (resp.choices[0].message.content or "").strip()
        except Exception as e:
            logger.warning("按页识图失败: %s", e)
            raise ValueError(f"按页识图失败: {str(e)}") from e

    try:
        cleaned = _clean_json_string(raw)
        data = json.loads(cleaned)
        if not isinstance(data, list):
            data = [data] if isinstance(data, dict) else []
        for item in data:
            if isinstance(item, dict):
                item["type"] = _normalize_question_type(item.get("type"), item.get("options"))
        data = _merge_same_number_questions(data)
        data = _merge_by_section_headers(data)
        return data
    except json.JSONDecodeError as e:
        logger.warning("按页识图返回非 JSON: %s", e)
        return []


# 裁剪附图最大边长，超过则等比缩小以减小体积
_CROP_IMAGE_MAX_SIDE = 800


def _crop_page_image_by_region(page_image_bytes: bytes, region: dict[str, float]) -> bytes | None:
    """
    按归一化区域 (x, y, width, height in 0-1) 从整页图中裁剪出该题对应区域，
    输出 JPEG（quality=85），超过最大边长则等比缩小，返回字节。失败返回 None。
    """
    try:
        img = Image.open(BytesIO(page_image_bytes))
        img = img.convert("RGB")
        W, H = img.size
        x = float(region.get("x", 0))
        y = float(region.get("y", 0))
        w = float(region.get("width", 1))
        h = float(region.get("height", 1))
        x1 = max(0, min(int(x * W), W - 1))
        y1 = max(0, min(int(y * H), H - 1))
        x2 = max(x1 + 1, min(int((x + w) * W), W))
        y2 = max(y1 + 1, min(int((y + h) * H), H))
        cropped = img.crop((x1, y1, x2, y2))
        cw, ch = cropped.size
        if max(cw, ch) > _CROP_IMAGE_MAX_SIDE:
            ratio = _CROP_IMAGE_MAX_SIDE / max(cw, ch)
            new_w = max(1, int(cw * ratio))
            new_h = max(1, int(ch * ratio))
            cropped = cropped.resize((new_w, new_h), Image.Resampling.LANCZOS)
        buf = BytesIO()
        cropped.save(buf, format="JPEG", quality=85)
        return buf.getvalue()
    except Exception as e:
        logger.warning("按区域裁剪页图失败: %s", e)
        return None


def attach_question_images_from_page(page_image_bytes: bytes, page_questions: list[dict[str, Any]]) -> None:
    """
    根据每题的 image_region / image_regions 从页图中裁剪出该题对应图片，写入 q["images"]（data URL 列表），
    并移除 image_region / image_regions 字段。无有效区域则 q["images"] = []。
    """
    for q in page_questions:
        if not isinstance(q, dict):
            continue
        regions: list[dict[str, float]] = []
        _append_region(regions, q.get("image_region"))
        for r in q.get("image_regions") or []:
            if isinstance(r, dict):
                regions.append(r)
        images: list[str] = []
        for reg in regions:
            cropped = _crop_page_image_by_region(page_image_bytes, reg)
            if cropped:
                b64 = base64.standard_b64encode(cropped).decode("ascii")
                images.append(f"data:image/jpeg;base64,{b64}")
        q["images"] = images
        q.pop("image_region", None)
        q.pop("image_regions", None)


def merge_page_questions(pages_questions: list[list[dict[str, Any]]]) -> list[dict[str, Any]]:
    """按页顺序合并题目列表；同题号的多段合并为一道，再统一题号（1, 2, 3, ...）。"""
    flat: list[dict[str, Any]] = []
    for page_list in pages_questions:
        for q in page_list:
            if isinstance(q, dict):
                flat.append(dict(q))
    flat = _merge_same_number_questions(flat)
    flat = _merge_by_section_headers(flat)
    for i, q in enumerate(flat, start=1):
        q["number"] = i
    return flat
