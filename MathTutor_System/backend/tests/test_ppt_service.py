"""
PPT 服务单测：_clean_json_string、create_pptx_file。
不调用真实 API。运行: 在 backend 目录下 python -m pytest tests/test_ppt_service.py -v
"""
import sys
from io import BytesIO

sys.path.insert(0, ".")

from app.services.ppt_service import _clean_json_string, create_pptx_file


def test_ppt_clean_json_string():
    """去除 markdown 代码块后得到纯 JSON 字符串。"""
    raw = """```json
{"title": "Lesson", "slides": [{"layout": "title", "title": "Hi"}]}
```"""
    out = _clean_json_string(raw)
    assert "```" not in out
    assert "Lesson" in out
    assert "slides" in out


def test_create_pptx_file_returns_bytes():
    """create_pptx_file 返回的 BytesIO 内含有效 .pptx 字节（PK 头）。"""
    content = {
        "title": "Test Lesson",
        "slides": [
            {"layout": "title", "title": "Main", "subtitle": "Sub"},
            {"layout": "content", "title": "Point 1", "bullets": ["A", "B", "C"]},
        ],
    }
    buf = create_pptx_file(content)
    assert isinstance(buf, BytesIO)
    data = buf.getvalue()
    assert len(data) > 500
    # .pptx 是 ZIP，文件头为 PK
    assert data[:2] == b"PK"


def test_create_pptx_file_empty_slides_title_only():
    """仅标题页时也能生成。"""
    content = {"title": "Only Title", "slides": [{"layout": "title", "title": "Only", "subtitle": ""}]}
    buf = create_pptx_file(content)
    assert len(buf.getvalue()) > 100
    assert buf.getvalue()[:2] == b"PK"


def test_create_pptx_file_content_no_bullets():
    """内容页无 bullets 时不报错。"""
    content = {
        "title": "T",
        "slides": [{"layout": "content", "title": "Empty body", "bullets": []}],
    }
    buf = create_pptx_file(content)
    assert len(buf.getvalue()) > 100


# ---------- PHASE 2C-5C：数学字段渲染集成 ----------

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def _pictures(slide):
    return [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]


def _text_shapes(slide):
    return [shape for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text.strip()]


def test_plain_deck_keeps_native_text_and_no_pictures():
    """纯文本字段必须保持原生 PowerPoint 文本，绝不栅格化。"""
    content = {
        "title": "纯文本讲义",
        "slides": [
            {"layout": "title", "title": " Plain Title ", "subtitle": "Sub"},
            {"layout": "content", "title": "学习目标", "bullets": ["理解直角三角形", "掌握判定"]},
        ],
    }
    prs = Presentation(create_pptx_file(content))
    assert len(prs.slides) == 2
    for slide in prs.slides:
        assert _pictures(slide) == []
    body_texts = " | ".join(shape.text_frame.text for shape in _text_shapes(prs.slides[1]))
    assert "理解直角三角形" in body_texts
    assert "学习目标" in body_texts


def test_math_fields_become_png_pictures():
    content = {
        "title": "勾股定理",
        "slides": [
            {"layout": "title", "title": "勾股定理 $a^2+b^2=c^2$", "subtitle": "学习 $c=5$ 的情形"},
            {
                "layout": "content",
                "title": "核心公式 $a^2+b^2=c^2$",
                "bullets": ["理解直角三角形", "当 $x=2$ 时，$y=x^2+1$ 的值为 $5$。", "$$c=\sqrt{a^2+b^2}$$"],
            },
        ],
    }
    prs = Presentation(create_pptx_file(content))
    title_pics = _pictures(prs.slides[0])
    assert len(title_pics) == 2  # 数学标题 + 数学副标题
    for pic in title_pics:
        assert pic.image.content_type == "image/png"
        assert pic.image.size[0] > 0 and pic.image.size[1] > 0
    content_pics = _pictures(prs.slides[1])
    assert len(content_pics) == 3  # 数学标题 + 混排 bullet + 展示公式 bullet
    body_texts = " | ".join(shape.text_frame.text for shape in _text_shapes(prs.slides[1]))
    assert "•  理解直角三角形" in body_texts  # 纯文本 bullet 仍是原生文本


def test_malformed_bullet_falls_back_to_raw_source_text():
    """畸形公式不得使 deck 失败，也不得静默丢弃：回退为规范源文本。"""
    content = {
        "title": "T",
        "slides": [
            {"layout": "content", "title": "回退", "bullets": ["$\\frac{1$"]},
        ],
    }
    prs = Presentation(create_pptx_file(content))
    assert _pictures(prs.slides[0]) == []
    body_texts = " | ".join(shape.text_frame.text for shape in _text_shapes(prs.slides[0]))
    assert "$\\frac{1$" in body_texts


def test_unsupported_math_command_falls_back_to_text():
    content = {
        "title": "T",
        "slides": [
            {"layout": "content", "title": "回退", "bullets": ["$\\href{javascript:alert(1)}{x}$"]},
        ],
    }
    prs = Presentation(create_pptx_file(content))
    assert _pictures(prs.slides[0]) == []
    body_texts = " | ".join(shape.text_frame.text for shape in _text_shapes(prs.slides[0]))
    assert "javascript:alert(1)" in body_texts


def test_hostile_text_bullet_never_becomes_picture_or_link():
    content = {
        "title": "T",
        "slides": [
            {"layout": "content", "title": "安全", "bullets": ["<script>alert(1)</script>"]},
        ],
    }
    prs = Presentation(create_pptx_file(content))
    assert _pictures(prs.slides[0]) == []
    body_texts = " | ".join(shape.text_frame.text for shape in _text_shapes(prs.slides[0]))
    assert "<script>alert(1)</script>" in body_texts


# ---------- RB01/RB02：真实 PPTX 几何不变量 ----------
#
# 权威下界：BODY_CONTENT_BOTTOM（由 _FOOTER_TOP / 卡片底边推导）。
# 对生成后的 PowerPoint 用 python-pptx 重开，直接断言每个正文 shape 的
# top >= body_top 且 top + height <= BODY_CONTENT_BOTTOM，并核对页脚不被覆盖。

from pptx.util import Inches

from app.services.ppt_service import (
    BODY_CONTENT_BOTTOM,
    _BODY_BOTTOM_PAD,
    _BODY_TOP,
    _CARD_BOTTOM,
    _CARD_BOTTOM_PAD,
    _EMU_PER_INCH,
    _FOOTER_SEPARATOR_LIFT,
    _FOOTER_TOP,
    _TITLE_BOX_H,
    _TITLE_TOP,
)

_GEOM_EPS = int(0.02 * _EMU_PER_INCH)  # 吸收放置时的整型截断
_TITLE_BOTTOM = _TITLE_TOP + _TITLE_BOX_H


def _content_slide_of(prs, index=0):
    """单内容页 deck 直接取该页；带 title 页的 deck 取第 2 页。"""
    return prs.slides[index] if len(prs.slides) == 1 else prs.slides[1]


def _card_top_of(slide):
    """内容卡 top：slide 上唯一的两块「自动形状 且 宽 >= 10in 且高 >= 3in」中的较上者。

    用 AUTO_SHAPE 类型排除 whole-body 单文本框（宽高也可能很大但是 TEXT_BOX）。
    """
    big = [
        s for s in slide.shapes
        if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
        and s.width >= Inches(10) and s.height >= Inches(3)
    ]
    assert len(big) == 2, f"expected card + shadow, got {len(big)}"
    return min(s.top for s in big)


def _assert_body_geometry(slide):
    """正文所有 flow shape 必须落在 [card_top+0.12, BODY_CONTENT_BOTTOM] 内。"""
    card_top = _card_top_of(slide)
    body_top_line = card_top + Inches(0.12)
    body = [
        s for s in slide.shapes
        if body_top_line - _GEOM_EPS <= s.top <= BODY_CONTENT_BOTTOM + _GEOM_EPS
    ]
    assert body, "content slide must have body flow shapes"
    for shape in body:
        assert shape.top >= body_top_line - _GEOM_EPS, f"body shape above body_top: {shape.shape_type}"
        assert shape.top + shape.height <= BODY_CONTENT_BOTTOM + _GEOM_EPS, (
            f"body shape crosses BODY_CONTENT_BOTTOM: {shape.shape_type} "
            f"bottom={_emu_in(shape.top + shape.height):.3f}in"
        )
    return body


def _assert_subtitle_band(slide):
    """副标题 shape 必须位于标题块之下、内容卡之上，且不与之重叠。"""
    card_top = _card_top_of(slide)
    band = [
        s for s in slide.shapes
        if _TITLE_BOTTOM - _GEOM_EPS <= s.top < card_top - _GEOM_EPS
    ]
    assert band, "expected subtitle shapes between title and card"
    for shape in band:
        assert shape.top >= _TITLE_BOTTOM - _GEOM_EPS
        assert shape.top + shape.height <= card_top + _GEOM_EPS, "subtitle overlaps content card"
    return band


def _bullet_text_shapes(slide):
    return [
        s for s in slide.shapes
        if s.has_text_frame and s.text_frame.text.strip().startswith("•")
    ]


def _emu_in(emu):
    return emu / _EMU_PER_INCH


def test_body_content_bottom_is_derived_from_footer_authority():
    """权威下界必须由页脚/卡片几何推导，且严格位于页脚分隔线之上。"""
    assert _CARD_BOTTOM == _FOOTER_TOP - _CARD_BOTTOM_PAD
    assert BODY_CONTENT_BOTTOM == _CARD_BOTTOM - _BODY_BOTTOM_PAD
    assert BODY_CONTENT_BOTTOM < _FOOTER_TOP - _FOOTER_SEPARATOR_LIFT


def test_rb01_scenario_a_short_bullets_with_math_bounded():
    """场景 A：12 条短要点（含数学）→ 所有正文 shape 都在正文可用区内。"""
    content = {
        "title": "T",
        "slides": [{
            "layout": "content",
            "title": "学习目标",
            "bullets": [f"直角三角形性质要点{i}" for i in range(11)] + ["勾股定理 $a^2+b^2=c^2$"],
        }],
    }
    prs = Presentation(create_pptx_file(content))
    slide = prs.slides[0]
    body = _assert_body_geometry(slide)
    # 无静默丢字段：11 条原生文本要点 + 1 张数学 PNG 全部在场。
    assert len(_bullet_text_shapes(slide)) == 11
    assert len(_pictures(slide)) == 1
    assert len(body) == 12


def test_rb01_scenario_b_twelve_math_bullets_bounded():
    """场景 B：12 条数学要点 → 12 张 PNG 全部在正文可用区内。"""
    content = {
        "title": "T",
        "slides": [{
            "layout": "content",
            "title": "公式训练",
            "bullets": [f"$x^{i}+{i}=0$" for i in range(12)],
        }],
    }
    prs = Presentation(create_pptx_file(content))
    slide = prs.slides[0]
    body = _assert_body_geometry(slide)
    assert len(_pictures(slide)) == 12
    assert len(body) == 12


def test_rb01_scenario_c_mixed_plain_math_fallback_bounded():
    """场景 C：plain / math / malformed 混排 → 全部落位，畸形回退源文本在场。"""
    bullets = []
    for i in range(4):
        bullets += [f"纯文本要点{i}", f"公式 $x^{i}$ 要点", "$\\frac{1$"]
    content = {
        "title": "T",
        "slides": [{"layout": "content", "title": "混合", "bullets": bullets}],
    }
    prs = Presentation(create_pptx_file(content))
    slide = prs.slides[0]
    _assert_body_geometry(slide)
    assert len(_pictures(slide)) == 4  # 4 条数学 PNG
    texts = " | ".join(s.text_frame.text for s in _bullet_text_shapes(slide))
    assert texts.count("纯文本要点") == 4
    assert texts.count("$\\frac{1$") == 4  # malformed 回退为源文本，不丢字段


def test_rb01_scenario_d_long_bullet_with_math_bounded():
    """场景 D：长要点 + 数学 → 收缩策略介入，但几何仍受权威下界约束。"""
    content = {
        "title": "T",
        "slides": [{
            "layout": "content",
            "title": "长文本与公式",
            "bullets": ["长" * 500, "勾股定理 $a^2+b^2=c^2$"],
        }],
    }
    prs = Presentation(create_pptx_file(content))
    slide = prs.slides[0]
    _assert_body_geometry(slide)
    assert len(_pictures(slide)) == 1
    texts = " | ".join(s.text_frame.text for s in _bullet_text_shapes(slide))
    assert "长" * 100 in texts  # 完整长文本在场（不截断）


def test_rb01_extreme_plain_density_whole_body_bounded_and_complete():
    """极端纯文本密度（12×500）：字号确定性下探，内容完整且不越界。"""
    content = {
        "title": "T",
        "slides": [{
            "layout": "content",
            "title": "密度极限",
            "bullets": ["超" * 500 for _ in range(12)],
        }],
    }
    prs = Presentation(create_pptx_file(content))
    slide = prs.slides[0]
    body = _assert_body_geometry(slide)
    boxes = [s for s in body if s.has_text_frame]
    assert len(boxes) == 1  # whole-body 单文本框
    paragraphs = [p.text for p in boxes[0].text_frame.paragraphs if p.text.strip()]
    assert len(paragraphs) == 12  # 12 条要点一条不少
    assert all(p == f"•  {'超' * 500}" for p in paragraphs)


def test_rb01_extreme_math_density_whole_body_fallback_no_silent_drop(caplog):
    """极端数学密度：whole-body safe fallback，数学条保留规范源文本，全部在场。"""
    import logging

    content = {
        "title": "T",
        "slides": [{
            "layout": "content",
            "title": "密度极限",
            "bullets": [f"压力测试{i} $a^2+b^2=c^2$ " + "长" * 500 for i in range(12)],
        }],
    }
    with caplog.at_level(logging.INFO, logger="app.services.ppt_math_renderer"):
        prs = Presentation(create_pptx_file(content))
    slide = prs.slides[0]
    _assert_body_geometry(slide)
    assert _pictures(slide) == []  # fallback 路径不放置任何正文 PNG
    boxes = [s for s in slide.shapes if s.has_text_frame and "压力测试0" in s.text_frame.text]
    assert len(boxes) == 1
    paragraphs = [p.text for p in boxes[0].text_frame.paragraphs if p.text.strip()]
    assert len(paragraphs) == 12
    assert all("$a^2+b^2=c^2$" in p for p in paragraphs)  # 数学以源文本保留
    assert any("body-density-fallback" in r.getMessage() for r in caplog.records)
    # 公式正文绝不进入日志
    assert "a^2+b^2=c^2" not in caplog.text
    assert "压力测试" not in caplog.text


def test_rb01_plain_path_still_bounded_without_math():
    """纯文本常规要点：单文本框必须落在正文可用区内（含 12 条要点）。"""
    content = {
        "title": "T",
        "slides": [{
            "layout": "content",
            "title": "常规",
            "bullets": [f"要点内容第{i}条，用于验证纯文本路径的几何边界。" for i in range(12)],
        }],
    }
    prs = Presentation(create_pptx_file(content))
    slide = prs.slides[0]
    _assert_body_geometry(slide)
    boxes = _bullet_text_shapes(slide)
    assert len(boxes) == 1  # 单文本框
    paragraphs = [p.text for p in boxes[0].text_frame.paragraphs if p.text.strip()]
    assert len(paragraphs) == 12  # 12 条要点全部在场


def test_rb02_plain_subtitle_survives_export():
    content = {
        "title": "T",
        "slides": [{
            "layout": "content",
            "title": "概念讲解",
            "subtitle": "直角三角形基础",
            "bullets": ["理解直角三角形"],
        }],
    }
    prs = Presentation(create_pptx_file(content))
    slide = prs.slides[0]
    texts = " | ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    assert "直角三角形基础" in texts  # 原生 PowerPoint 文本
    band = _assert_subtitle_band(slide)
    subtitle_boxes = [s for s in band if s.has_text_frame and "直角三角形基础" in s.text_frame.text]
    assert len(subtitle_boxes) == 1
    _assert_body_geometry(slide)


def test_rb02_math_subtitle_becomes_picture():
    content = {
        "title": "T",
        "slides": [{
            "layout": "content",
            "title": "概念讲解",
            "subtitle": "学习 $c=5$ 的情形",
            "bullets": ["理解直角三角形"],
        }],
    }
    prs = Presentation(create_pptx_file(content))
    slide = prs.slides[0]
    band = _assert_subtitle_band(slide)
    subtitle_pics = [s for s in band if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(subtitle_pics) == 1
    assert subtitle_pics[0].image.content_type == "image/png"
    assert _pictures(slide) == subtitle_pics  # body 无数学 → 无正文 PNG
    _assert_body_geometry(slide)


def test_rb02_malformed_subtitle_survives_as_raw_text():
    content = {
        "title": "T",
        "slides": [{
            "layout": "content",
            "title": "概念讲解",
            "subtitle": "$\\frac{1$",
            "bullets": ["理解直角三角形"],
        }],
    }
    prs = Presentation(create_pptx_file(content))
    slide = prs.slides[0]
    texts = " | ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    assert "$\\frac{1$" in texts  # renderer 失败 → 原样源文本，不丢
    assert _pictures(slide) == []
    _assert_subtitle_band(slide)
    _assert_body_geometry(slide)


def test_rb02_subtitle_with_max_body_stays_bounded():
    """副标题 + 满载正文：副标题带、正文区、页脚互不越界。"""
    content = {
        "title": "T",
        "slides": [{
            "layout": "content",
            "title": "满载页 $x^2$",
            "subtitle": "本页要点较多，请结合例题理解每个公式的适用条件与推导过程。" * 3,
            "bullets": [f"公式 $x^{i}+{i}=0$ 与应用要点说明" for i in range(12)],
        }],
    }
    prs = Presentation(create_pptx_file(content))
    slide = prs.slides[0]
    _assert_subtitle_band(slide)
    body = _assert_body_geometry(slide)
    assert len(body) == 12  # 12 条数学要点全部在场
    # 页脚分隔线之上无正文越界（权威下界严格小于分隔线 top）
    assert BODY_CONTENT_BOTTOM < _FOOTER_TOP - _FOOTER_SEPARATOR_LIFT


def test_rb02_no_subtitle_keeps_legacy_body_top():
    """无副标题内容页保持既有 body_top 几何，回归保护。"""
    content = {
        "title": "T",
        "slides": [{"layout": "content", "title": "常规页", "bullets": ["要点"]}],
    }
    prs = Presentation(create_pptx_file(content))
    slide = prs.slides[0]
    card_top = _card_top_of(slide)
    assert card_top == _BODY_TOP - Inches(0.12)
    _assert_body_geometry(slide)


if __name__ == "__main__":
    test_ppt_clean_json_string()
    test_create_pptx_file_returns_bytes()
    test_create_pptx_file_empty_slides_title_only()
    test_create_pptx_file_content_no_bullets()
    print("test_ppt_service: 全部通过")
