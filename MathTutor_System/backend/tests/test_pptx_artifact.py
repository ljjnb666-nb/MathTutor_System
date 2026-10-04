"""PHASE 2C-5D：Magic PPTX 工件结构验证。

不止 PK 头：真实 deck 必须能用 Presentation 重开、ZIP 完整性通过、
数学页有 PNG 媒体关系、纯文本页保持原生文本、畸形回退 deck 可重开、
slide XML 中不存在可执行标记（script / 超链接关系）。
"""
import io
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.services.ppt_math_renderer import (
    MAX_RENDER_HEIGHT_PX,
    MAX_RENDER_PIXELS_PER_IMAGE,
    MAX_RENDER_WIDTH_PX,
)
from app.services.ppt_service import create_pptx_file

MIXED_DECK = {
    "title": "勾股定理 $a^2+b^2=c^2$",
    "slides": [
        {"layout": "title", "title": "勾股定理 $a^2+b^2=c^2$", "subtitle": "学习 $c=5$ 的情形"},
        {
            "layout": "content",
            "title": "核心公式 $a^2+b^2=c^2$",
            "bullets": [
                "理解直角三角形",
                "当 $x=2$ 时，$y=x^2+1$ 的值为 $5$。",
                "$$c=\\sqrt{a^2+b^2}$$",
                "$\\frac{1$",
                "$\\href{javascript:alert(1)}{x}$",
                "<script>alert(1)</script>",
            ],
        },
        {"layout": "content", "title": "本课总结", "bullets": ["全部原生文本", "没有任何公式"]},
    ],
}


def _pictures(slide):
    return [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]


def _media_names(zf):
    return sorted(n for n in zf.namelist() if n.startswith("ppt/media/"))


def test_mixed_deck_reopens_with_expected_structure():
    data = create_pptx_file(MIXED_DECK).getvalue()

    assert data[:2] == b"PK"
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 3

    # ZIP 完整性与 OOXML 结构
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        assert zf.testzip() is None
        names = zf.namelist()
        slide_xmls = sorted(n for n in names if n.startswith("ppt/slides/slide") and n.endswith(".xml"))
        assert len(slide_xmls) == 3
        assert any(n.startswith("ppt/slides/_rels/") for n in names)
        assert "[Content_Types].xml" in names
        media = _media_names(zf)
        assert media, "math deck must embed rendered media"


def test_math_slides_have_png_media_relationships():
    data = create_pptx_file(MIXED_DECK).getvalue()
    prs = Presentation(io.BytesIO(data))

    # 标题页：数学主标题 + 数学副标题 → 2 张 PNG
    title_pics = _pictures(prs.slides[0])
    assert len(title_pics) == 2
    # 内容页：数学标题 + 混排 bullet + 展示公式 bullet → 3 张 PNG
    content_pics = _pictures(prs.slides[1])
    assert len(content_pics) == 3

    for pic in title_pics + content_pics:
        blob = pic.image.blob
        assert blob[:8] == b"\x89PNG\r\n\x1a\n"
        with Image.open(io.BytesIO(blob)) as im:
            assert im.format == "PNG"
            width_px, height_px = im.size
        assert 0 < width_px <= MAX_RENDER_WIDTH_PX
        assert 0 < height_px <= MAX_RENDER_HEIGHT_PX
        assert width_px * height_px <= MAX_RENDER_PIXELS_PER_IMAGE

    # ZIP 级别：媒体全部是 PNG，且被 slide 关系引用
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for name in _media_names(zf):
            assert zf.read(name)[:8] == b"\x89PNG\r\n\x1a\n"
        slide2_rels = zf.read("ppt/slides/_rels/slide2.xml.rels").decode("utf-8")
        assert "../media/image" in slide2_rels


def test_plain_slide_keeps_real_text_without_media():
    data = create_pptx_file(MIXED_DECK).getvalue()
    prs = Presentation(io.BytesIO(data))

    slide3 = prs.slides[2]
    assert _pictures(slide3) == []
    body_texts = " | ".join(
        shape.text_frame.text for shape in slide3.shapes if shape.has_text_frame
    )
    assert "全部原生文本" in body_texts
    assert "本课总结" in body_texts

    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        slide3_rels = zf.read("ppt/slides/_rels/slide3.xml.rels").decode("utf-8")
        assert "ppt/media/" not in slide3_rels


def test_malformed_fallback_deck_reopens():
    content = {
        "title": "T",
        "slides": [
            {"layout": "title", "title": "畸形 $\\frac{1$", "subtitle": "$\\sqrt"},
            {"layout": "content", "title": "回退 $\\href{x}{y}$", "bullets": ["$\\frac{1$", "正常 $x$ 公式"]},
        ],
    }
    data = create_pptx_file(content).getvalue()
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 2
    slide2_texts = " | ".join(
        shape.text_frame.text for shape in prs.slides[1].shapes if shape.has_text_frame
    )
    # 畸形 bullet 与不受支持命令的标题都保留规范源文本，不静默丢弃。
    assert "$\\frac{1$" in slide2_texts
    assert "回退 $\\href{x}{y}$" in slide2_texts
    # 同一页的合法公式 bullet 仍然渲染为 PNG（混合回退互不影响）。
    assert len(_pictures(prs.slides[1])) == 1


def test_slide_xml_contains_no_executable_markup():
    data = create_pptx_file(MIXED_DECK).getvalue()
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for name in zf.namelist():
            if not (name.startswith("ppt/slides/slide") and name.endswith(".xml")):
                continue
            xml = zf.read(name).decode("utf-8")
            assert "<script" not in xml.lower()
            assert "<svg" not in xml.lower()
            # 数学源绝不生成超链接；敌对文本只能以转义文本形态存在
            assert "hlinkClick" not in xml
            for rels_name in ("ppt/_rels/presentation.xml.rels",):
                rels = zf.read(rels_name).decode("utf-8")
                assert 'Target="javascript:' not in rels
        for name in zf.namelist():
            if name.endswith(".rels"):
                assert 'Target="javascript:' not in zf.read(name).decode("utf-8")
