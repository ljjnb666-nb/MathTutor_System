"""PHASE 2C-5D：/api/tools/generate-ppt 与 /api/tools/build-pptx 端点契约。

- 不调用真实 LLM：generate-ppt 在服务边界打 mock；build-pptx 走真实
  渲染器与真实构建器；
- 保持鉴权、magic_ppt 套餐门禁、受控错误、内容上限、文件名净化。
"""
import io
from urllib.parse import unquote

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.api.endpoints import tools as tools_endpoint
from app.api.endpoints.auth import get_current_user
from app.core.deps import LLMConfig, get_llm_config
from app.main import app
from app.models.user import User
from app.schemas.ppt_dto import PPTContent


@pytest.fixture
def teacher_client():
    user = User(id=1, username="ppt-teacher", is_active=True, role="teacher")
    app.dependency_overrides[get_current_user] = lambda: user
    app.dependency_overrides[get_llm_config] = lambda: LLMConfig(
        provider="deepseek", api_key="k", base_url="https://api.deepseek.com", model="m",
    )
    client = TestClient(app)
    yield client
    app.dependency_overrides.pop(get_current_user, None)
    app.dependency_overrides.pop(get_llm_config, None)


def _allow_subscription(monkeypatch):
    """放行套餐门禁并记录调用（不触库）。"""
    calls = {"feature": []}

    def fake_get_subscription(current_user, db):
        return object()

    def fake_require_feature(subscription, feature_key, current_user, db=None):
        calls["feature"].append(feature_key)

    monkeypatch.setattr(tools_endpoint, "get_current_subscription", fake_get_subscription)
    monkeypatch.setattr(tools_endpoint, "require_feature", fake_require_feature)
    return calls


def test_generate_ppt_requires_auth():
    client = TestClient(app)
    response = client.post("/api/tools/generate-ppt", json={"topic": "勾股定理"})
    assert response.status_code == 401


def test_build_pptx_requires_auth():
    client = TestClient(app)
    response = client.post("/api/tools/build-pptx", json={"slides": [{"layout": "content"}]})
    assert response.status_code == 401


def test_generate_ppt_enforces_magic_ppt_gate(monkeypatch, teacher_client):
    calls = {"feature": []}

    def fake_get_subscription(current_user, db):
        return object()

    def fake_require_feature(subscription, feature_key, current_user, db=None):
        calls["feature"].append(feature_key)
        raise HTTPException(status_code=403, detail="MAGIC_PPT_REQUIRES_PRO")

    monkeypatch.setattr(tools_endpoint, "get_current_subscription", fake_get_subscription)
    monkeypatch.setattr(tools_endpoint, "require_feature", fake_require_feature)

    response = teacher_client.post("/api/tools/generate-ppt", json={"topic": "勾股定理"})
    assert response.status_code == 403
    assert calls["feature"] == ["magic_ppt"]


def test_build_pptx_enforces_magic_ppt_gate(monkeypatch, teacher_client):
    def fake_require_feature(subscription, feature_key, current_user, db=None):
        raise HTTPException(status_code=403, detail="MAGIC_PPT_REQUIRES_PRO")

    monkeypatch.setattr(tools_endpoint, "get_current_subscription", lambda current_user, db: object())
    monkeypatch.setattr(tools_endpoint, "require_feature", fake_require_feature)

    response = teacher_client.post(
        "/api/tools/build-pptx", json={"slides": [{"layout": "content", "title": "T"}]}
    )
    assert response.status_code == 403


def test_generate_ppt_returns_canonical_content(monkeypatch, teacher_client):
    _allow_subscription(monkeypatch)
    captured = {}

    def fake_generate(topic, grade, **kwargs):
        captured["topic"] = topic
        captured["grade"] = grade
        return PPTContent.model_validate({
            "title": "勾股定理",
            "slides": [
                {"layout": "title", "title": "勾股定理", "subtitle": "直角三角形"},
                {"layout": "content", "title": "核心公式", "bullets": ["$a^2+b^2=c^2$"]},
            ],
        })

    monkeypatch.setattr(tools_endpoint, "generate_lecture_content", fake_generate)

    response = teacher_client.post("/api/tools/generate-ppt", json={"topic": "勾股定理", "grade": "Middle"})

    assert response.status_code == 200
    assert captured == {"topic": "勾股定理", "grade": "Middle"}
    data = response.json()
    assert set(data.keys()) == {"title", "slides"}
    assert data["slides"][1] == {
        "layout": "content", "title": "核心公式", "subtitle": "", "bullets": ["$a^2+b^2=c^2$"],
    }


def test_generate_ppt_invalid_llm_response_is_controlled_400(monkeypatch, teacher_client):
    _allow_subscription(monkeypatch)

    def fake_generate(topic, grade, **kwargs):
        raise ValueError("LLM_RESPONSE_INVALID: 讲稿服务返回格式无法解析，请调整主题后重试。")

    monkeypatch.setattr(tools_endpoint, "generate_lecture_content", fake_generate)

    response = teacher_client.post("/api/tools/generate-ppt", json={"topic": "勾股定理"})
    assert response.status_code == 400
    assert response.json()["detail"].startswith("LLM_RESPONSE_INVALID")


def test_generate_ppt_provider_error_is_controlled_400(monkeypatch, teacher_client):
    _allow_subscription(monkeypatch)

    def fake_generate(topic, grade, **kwargs):
        raise ValueError("LLM_PROVIDER_ERROR: 讲稿生成服务暂时不可用，请稍后重试。")

    monkeypatch.setattr(tools_endpoint, "generate_lecture_content", fake_generate)

    response = teacher_client.post("/api/tools/generate-ppt", json={"topic": "勾股定理"})
    assert response.status_code == 400
    assert response.json()["detail"].startswith("LLM_PROVIDER_ERROR")


def test_generate_ppt_unexpected_error_is_safe_500(monkeypatch, teacher_client):
    _allow_subscription(monkeypatch)

    def fake_generate(topic, grade, **kwargs):
        raise RuntimeError("internal socket details with key=sk-secret12345")

    monkeypatch.setattr(tools_endpoint, "generate_lecture_content", fake_generate)

    response = teacher_client.post("/api/tools/generate-ppt", json={"topic": "勾股定理"})
    assert response.status_code == 500
    detail = response.json()["detail"]
    assert detail.startswith("PPT_GENERATION_ERROR")
    assert "sk-secret12345" not in response.text
    assert "RuntimeError" not in detail


def test_generate_ppt_requires_topic(monkeypatch, teacher_client):
    _allow_subscription(monkeypatch)
    response = teacher_client.post("/api/tools/generate-ppt", json={"topic": "   "})
    assert response.status_code == 400


def test_build_pptx_with_real_renderer_and_builder(monkeypatch, teacher_client):
    _allow_subscription(monkeypatch)
    body = {
        "title": "勾股定理",
        "filename": "勾股 定理-01",
        "slides": [
            {"layout": "title", "title": "勾股定理 $a^2+b^2=c^2$", "subtitle": "学习 $c=5$"},
            {"layout": "content", "title": "总结", "bullets": ["纯文本要点", "公式 $x^2$ 要点"]},
        ],
    }
    response = teacher_client.post("/api/tools/build-pptx", json=body)

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    disposition = response.headers["content-disposition"]
    # RFC 5987：中文文件名经 filename*=UTF-8'' 传递，ASCII 兜底保证头可编码。
    assert disposition.startswith('attachment; filename="Lesson.pptx"')
    star_value = disposition.split("filename*=UTF-8''")[1]
    assert unquote(star_value) == "勾股 定理-01.pptx"
    data = response.content
    assert data[:2] == b"PK"
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 2
    pics = [s for s in prs.slides[1].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pics) == 1  # 只有数学 bullet 渲染为图片，纯文本要点保持文本


def test_build_pptx_ascii_filename_kept_verbatim(monkeypatch, teacher_client):
    _allow_subscription(monkeypatch)
    response = teacher_client.post(
        "/api/tools/build-pptx",
        json={"filename": "Lesson One-01", "slides": [{"layout": "content", "title": "T"}]},
    )
    assert response.status_code == 200
    disposition = response.headers["content-disposition"]
    assert disposition.startswith('attachment; filename="Lesson One-01.pptx"')
    assert unquote(disposition.split("filename*=UTF-8''")[1]) == "Lesson One-01.pptx"


def test_build_pptx_filename_sanitized(monkeypatch, teacher_client):
    _allow_subscription(monkeypatch)
    response = teacher_client.post(
        "/api/tools/build-pptx",
        json={"filename": "../../evil name<script>", "slides": [{"layout": "content", "title": "T"}]},
    )
    assert response.status_code == 200
    disposition = response.headers["content-disposition"]
    assert ".." not in disposition
    plain_name = disposition.split('filename="')[1].split('"')[0]
    assert "/" not in plain_name
    assert unquote(disposition.split("filename*=UTF-8''")[1]) == "evil namescript.pptx"


def test_build_pptx_default_filename(monkeypatch, teacher_client):
    _allow_subscription(monkeypatch)
    response = teacher_client.post(
        "/api/tools/build-pptx", json={"slides": [{"layout": "content", "title": "T"}]}
    )
    assert response.status_code == 200
    assert 'filename="Lesson.pptx"' in response.headers["content-disposition"]


def test_build_pptx_enforces_content_limits(monkeypatch, teacher_client):
    _allow_subscription(monkeypatch)

    too_many_slides = {"slides": [{"layout": "content", "title": f"S{i}"} for i in range(31)]}
    assert teacher_client.post("/api/tools/build-pptx", json=too_many_slides).status_code == 422

    over_bullet = {
        "slides": [{"layout": "content", "title": "T", "bullets": ["x" * 501]}],
    }
    assert teacher_client.post("/api/tools/build-pptx", json=over_bullet).status_code == 422

    boundary = {
        "slides": [{"layout": "content", "title": "T", "bullets": ["x" * 500] * 12}],
    }
    assert teacher_client.post("/api/tools/build-pptx", json=boundary).status_code == 200

    empty = {"slides": []}
    assert teacher_client.post("/api/tools/build-pptx", json=empty).status_code == 422


# ---------- RB05：canonical layout enum（build 端点拒绝非法 layout，不静默丢弃） ----------


@pytest.mark.parametrize("layout", ["section", "unknown", "Title"])
def test_build_pptx_rejects_non_canonical_layout(monkeypatch, teacher_client, layout):
    _allow_subscription(monkeypatch)
    response = teacher_client.post(
        "/api/tools/build-pptx",
        json={"slides": [{"layout": layout, "title": "T"}]},
    )
    assert response.status_code == 422


@pytest.mark.parametrize("layout", ["title", "content"])
def test_build_pptx_accepts_canonical_layout(monkeypatch, teacher_client, layout):
    _allow_subscription(monkeypatch)
    response = teacher_client.post(
        "/api/tools/build-pptx",
        json={"slides": [{"layout": layout, "title": "T"}]},
    )
    assert response.status_code == 200


# ---------- RB01-EXT2：物理不可渲染 payload → 受控 PPT_LAYOUT_UNFIT ----------


def test_build_pptx_body_impossible_fit_controlled_422(monkeypatch, teacher_client):
    _allow_subscription(monkeypatch)
    dense = "\n".join(["x"] * 250)  # 499 字符 schema 合法，但 12×250 行物理放不下
    response = teacher_client.post(
        "/api/tools/build-pptx",
        json={"slides": [{"layout": "content", "title": "T", "bullets": [dense] * 12}]},
    )
    assert response.status_code == 422
    detail = response.json()["detail"]
    assert detail.startswith("PPT_LAYOUT_UNFIT")
    assert "当前幻灯片内容过多" in detail
    # 不泄漏原始输入、堆栈或异常类
    assert "x" * 50 not in response.text
    assert "Traceback" not in response.text
    assert "PPTLayoutUnfitError" not in response.text


def test_build_pptx_title_slide_subtitle_unfit_controlled_422(monkeypatch, teacher_client):
    _allow_subscription(monkeypatch)
    dense = "\n".join(["x"] * 250)
    response = teacher_client.post(
        "/api/tools/build-pptx",
        json={"slides": [{"layout": "title", "title": "主标题", "subtitle": dense}]},
    )
    assert response.status_code == 422
    assert response.json()["detail"].startswith("PPT_LAYOUT_UNFIT")
    assert "x" * 50 not in response.text
