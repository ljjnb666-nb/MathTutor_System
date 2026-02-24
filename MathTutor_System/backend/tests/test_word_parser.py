"""
Word 试卷解析服务单测：_clean_json_string、extract_text_from_docx。
不调用真实 API。运行: 在 backend 目录下 python -m pytest tests/test_word_parser.py -v
"""
import sys
from io import BytesIO

sys.path.insert(0, ".")

from docx import Document

from app.services.word_parser import _clean_json_string, _normalize_question_type, extract_text_from_docx


def test_clean_json_string_plain():
    """无 markdown 包裹时原样返回（去除首尾空白）。"""
    raw = '  {"questions": [{"content": "1+1?", "type": "choice"}]}  '
    out = _clean_json_string(raw)
    assert out.strip() == raw.strip()
    assert "questions" in out


def test_clean_json_string_with_markdown_fence():
    """带 ```json ... ``` 时提取中间部分。"""
    raw = """```json
{"title": "Exam", "slides": []}
```"""
    out = _clean_json_string(raw)
    assert "```" not in out
    assert "title" in out
    assert "Exam" in out


def test_clean_json_string_fence_no_json_tag():
    """仅 ``` ... ``` 也提取中间部分。"""
    raw = """```
[{"type": "choice", "content": "题1"}]
```"""
    out = _clean_json_string(raw)
    assert "```" not in out
    assert "choice" in out


def test_clean_json_string_empty():
    """空字符串或仅空白返回空。"""
    assert _clean_json_string("") == ""
    assert _clean_json_string("   ") == ""


def test_normalize_question_type():
    """题型映射：英文、中文、别名均映射为 choice | fill | solution；未知+有选项推断为 choice。"""
    assert _normalize_question_type("choice") == "choice"
    assert _normalize_question_type("选择题") == "choice"
    assert _normalize_question_type("单选") == "choice"
    assert _normalize_question_type("判断题") == "choice"
    assert _normalize_question_type("fill") == "fill"
    assert _normalize_question_type("填空题") == "fill"
    assert _normalize_question_type("text") == "fill"
    assert _normalize_question_type("solution") == "solution"
    assert _normalize_question_type("解答题") == "solution"
    assert _normalize_question_type("计算题") == "solution"
    assert _normalize_question_type("应用题") == "solution"
    assert _normalize_question_type("unknown", []) == "fill"
    assert _normalize_question_type("", ["A.1", "B.2"]) == "choice"
    assert _normalize_question_type(None, ["A", "B", "C"]) == "choice"


def test_extract_text_from_docx():
    """从内存中的 .docx 提取段落文本。"""
    doc = Document()
    doc.add_paragraph("第一题：1+1=?")
    doc.add_paragraph("第二题：勾股定理的应用")
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)

    text = extract_text_from_docx(buf)
    assert "第一题" in text
    assert "1+1" in text
    assert "第二题" in text
    assert "勾股定理" in text


if __name__ == "__main__":
    test_clean_json_string_plain()
    test_clean_json_string_with_markdown_fence()
    test_clean_json_string_fence_no_json_tag()
    test_clean_json_string_empty()
    test_normalize_question_type()
    test_extract_text_from_docx()
    print("test_word_parser: 全部通过")
